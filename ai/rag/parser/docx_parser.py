"""DOCX 텍스트 추출 (SRS PAR-01)."""
import io
from typing import Dict, List


def extract_docx(source) -> List[Dict]:
    """DOCX에서 문단 단위 텍스트를 추출한다.

    source: 파일 경로(str) 또는 파일 내용(bytes)
    반환: [{"page": None, "para": 1, "text": "..."}]
    """
    from docx import Document

    # 워드 파일에는 페이지 구분이 저장되어 있지 않다(화면·인쇄 설정에 따라 달라진다).
    # 그래서 page 대신 문단 번호를 위치 정보로 사용한다. 빈 문단은 제외한다.
    doc = Document(io.BytesIO(source)) if isinstance(source, (bytes, bytearray)) else Document(source)
    return [
        {"page": None, "para": i, "text": p.text}
        for i, p in enumerate(doc.paragraphs, start=1)
        if p.text.strip()
    ]
